"""
Builds the Excel workbook you hand back to a distributor after running
reconcile() — this is the V0 "manual service" deliverable itself.

Summary sheet numbers are formulas that reference the other sheets (per the
project convention: never hardcode a computed total), so the workbook stays
correct if you manually move a row between sheets during review.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
LABEL_FONT = Font(name=FONT_NAME, bold=True)
BODY_FONT = Font(name=FONT_NAME)
CURRENCY_FMT = '#,##0.00 "K"'
DATE_FMT = "yyyy-mm-dd"
MAX_DATA_ROWS = 1000  # summary formula ranges assume no sheet exceeds this


def _write_sheet(ws, columns: list[tuple[str, str]], df: pd.DataFrame):
    """`columns` is a list of (header_label, dataframe_column) pairs, in
    the order they should appear. Missing dataframe columns are left blank."""
    for c, (label, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for r, (_, row) in enumerate(df.iterrows(), start=2):
        for c, (_, field) in enumerate(columns, start=1):
            value = row.get(field, "")
            if pd.isna(value):
                value = ""
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            if field in ("amount", "invoice_total", "remaining_balance",
                         "invoice_balance", "balance"):
                cell.number_format = CURRENCY_FMT
            if field == "date":
                cell.number_format = DATE_FMT

    for c, (label, _) in enumerate(columns, start=1):
        width = max(14, len(label) + 2)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"


def write_report(results: dict[str, pd.DataFrame], output_path: str | Path,
                  business_name: str = ""):
    wb = Workbook()
    wb.remove(wb.active)

    txn_cols = [
        ("Transaction ID", "transaction_id"), ("Date", "date"), ("Amount", "amount"),
        ("Sender", "sender_name"), ("Sender Phone", "sender_phone"),
        ("Reference", "reference"), ("Matched Invoice", "matched_invoice"),
        ("Customer", "customer"), ("Match Rule", "match_rule"),
    ]
    partial_cols = txn_cols + [("Invoice Total", "invoice_total"),
                                ("Remaining Balance", "remaining_balance")]
    review_cols = txn_cols + [("Details", "candidate_invoices")]
    unmatched_cols = [
        ("Transaction ID", "transaction_id"), ("Date", "date"), ("Amount", "amount"),
        ("Sender", "sender_name"), ("Sender Phone", "sender_phone"),
        ("Reference", "reference"),
    ]
    open_inv_cols = [
        ("Invoice ID", "invoice_id"), ("Customer", "customer_name"),
        ("Phone", "customer_phone"), ("Invoice Date", "date"),
        ("Invoice Amount", "amount"), ("Balance Outstanding", "balance"),
    ]

    _write_sheet(wb.create_sheet("Matched"), txn_cols, results["matched"])
    _write_sheet(wb.create_sheet("Partial"), partial_cols, results["partial"])
    _write_sheet(wb.create_sheet("Needs Review"), review_cols, results["needs_review"])
    _write_sheet(wb.create_sheet("Unmatched"), unmatched_cols, results["unmatched"])
    _write_sheet(wb.create_sheet("Open Invoices"), open_inv_cols, results["open_invoices"])

    # "Open Invoices" gets a Days Outstanding formula column (col G)
    ws_open = wb["Open Invoices"]
    ws_open.cell(row=1, column=7, value="Days Outstanding").font = HEADER_FONT
    ws_open.cell(row=1, column=7).fill = HEADER_FILL
    for r in range(2, len(results["open_invoices"]) + 2):
        ws_open.cell(row=r, column=7, value=f"=TODAY()-D{r}").font = BODY_FONT
    ws_open.column_dimensions["G"].width = 16

    _write_summary(wb.create_sheet("Summary", 0), business_name)
    wb.save(output_path)


def _write_summary(ws, business_name: str):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18

    title = f"Reconciliation Summary — {business_name}" if business_name else "Reconciliation Summary"
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Generated:"
    ws["A2"].font = LABEL_FONT
    ws["B2"] = "=TODAY()"
    ws["B2"].number_format = DATE_FMT
    ws["B2"].font = BODY_FONT

    rng = lambda sheet, col: f"'{sheet}'!{col}2:{col}{MAX_DATA_ROWS}"

    rows = [
        ("Transactions imported", "count",
         f"=COUNTA({rng('Matched','A')})+COUNTA({rng('Partial','A')})+"
         f"COUNTA({rng('Needs Review','A')})+COUNTA({rng('Unmatched','A')})"),
        ("Total amount received", "money",
         f"=SUM({rng('Matched','C')})+SUM({rng('Partial','C')})+"
         f"SUM({rng('Needs Review','C')})+SUM({rng('Unmatched','C')})"),
        ("", "", ""),
        ("Automatically matched — count", "count", f"=COUNTA({rng('Matched','A')})"),
        ("Automatically matched — value", "money", f"=SUM({rng('Matched','C')})"),
        ("Partial payments — count", "count", f"=COUNTA({rng('Partial','A')})"),
        ("Partial payments — value", "money", f"=SUM({rng('Partial','C')})"),
        ("Needs review — count", "count", f"=COUNTA({rng('Needs Review','A')})"),
        ("Needs review — value", "money", f"=SUM({rng('Needs Review','C')})"),
        ("Unmatched — count", "count", f"=COUNTA({rng('Unmatched','A')})"),
        ("Unmatched — value", "money", f"=SUM({rng('Unmatched','C')})"),
        ("", "", ""),
        ("Open invoices remaining — count", "count", f"=COUNTA({rng('Open Invoices','A')})"),
        ("Open invoices — outstanding value", "money", f"=SUM({rng('Open Invoices','F')})"),
        ("", "", ""),
        ("Auto-match rate", "pct", "=B5/B2"),
    ]

    r = 4
    matched_count_row = None
    imported_row = None
    for label, kind, formula in rows:
        ws.cell(row=r, column=1, value=label).font = (LABEL_FONT if label else BODY_FONT)
        if formula:
            cell = ws.cell(row=r, column=2, value=formula)
            cell.font = BODY_FONT
            if kind == "money":
                cell.number_format = CURRENCY_FMT
            elif kind == "pct":
                cell.number_format = "0.0%"
        if label == "Transactions imported":
            imported_row = r
        if label == "Automatically matched — count":
            matched_count_row = r
        r += 1

    # Fix the match-rate formula to the actual rows used above (rows shift if
    # this list is edited later, so compute the reference rather than hardcode).
    if matched_count_row and imported_row:
        ws.cell(row=r - 1, column=2,
                value=f"=B{matched_count_row}/B{imported_row}")

    note = ("Note: 'Auto-match rate' only counts the Matched sheet as fully "
            "resolved. Partial and Needs Review still require a human look — "
            "that split is the actual baseline this tool should be judged on.")
    ws.cell(row=r + 1, column=1, value=note).font = Font(name=FONT_NAME, italic=True, size=9)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=6)


def write_aging_report(aging_df: pd.DataFrame, summary_df: pd.DataFrame,
                        output_path: str | Path, business_name: str = ""):
    """The accounts-receivable view: 'who owes me money, and how overdue is
    it' — read from the persisted invoice balances in db.py, not from a
    single reconcile() run, so it reflects everything accumulated so far."""
    wb = Workbook()
    wb.remove(wb.active)

    title = f"Accounts Receivable Aging — {business_name}" if business_name else "Accounts Receivable Aging"
    ws_sum = wb.create_sheet("By Customer", 0)
    ws_sum.sheet_view.showGridLines = False
    ws_sum["A1"] = title
    ws_sum["A1"].font = TITLE_FONT
    ws_sum["A2"] = "As of:"
    ws_sum["A2"].font = LABEL_FONT
    ws_sum["B2"] = "=TODAY()"
    ws_sum["B2"].number_format = DATE_FMT
    ws_sum["B2"].font = BODY_FONT

    bucket_cols = [c for c in summary_df.columns if c not in ("customer_name", "total_owed")]
    summary_cols = [("Customer", "customer_name"), ("Total Owed", "total_owed")] + \
                   [(f"{b} days", b) for b in bucket_cols]
    for c, (label, _) in enumerate(summary_cols, start=1):
        cell = ws_sum.cell(row=4, column=c, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r, (_, row) in enumerate(summary_df.iterrows(), start=5):
        for c, (_, field) in enumerate(summary_cols, start=1):
            cell = ws_sum.cell(row=r, column=c, value=row.get(field, ""))
            cell.font = BODY_FONT
            if field != "customer_name":
                cell.number_format = CURRENCY_FMT
    for c, (label, _) in enumerate(summary_cols, start=1):
        ws_sum.column_dimensions[get_column_letter(c)].width = max(16, len(label) + 2)
    ws_sum.freeze_panes = "A5"

    detail_cols = [
        ("Invoice ID", "invoice_id"), ("Customer", "customer_name"),
        ("Phone", "customer_phone"), ("Invoice Date", "date"),
        ("Invoice Amount", "amount"), ("Balance Outstanding", "balance"),
        ("Days Outstanding", "days_outstanding"), ("Bucket", "bucket"),
    ]
    _write_sheet(wb.create_sheet("By Invoice"), detail_cols, aging_df)

    wb.save(output_path)
